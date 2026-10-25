#!/usr/bin/env python3
"""extract_docs.py - 提取目录下所有 xlsx/docx 的文本摘要, 生成 _INDEX.md 索引。

用法:
  python3 extract_docs.py <目录> [--out _INDEX.md]

只取每份文档前若干行/段落看结构(避免爆上下文), 不提取图片。
输出 _INDEX.md 供 agent 后续检索引用。
"""
import os
import sys
import glob
import argparse
from openpyxl import load_workbook
from docx import Document


def clean(s):
    return (s or "").strip().replace("\n", " ").replace("|", "/")


def extract_xlsx(path, max_rows=8):
    wb = load_workbook(path, read_only=True, data_only=True)
    parts = []
    for ws in wb.worksheets:
        parts.append(f"\n#### Sheet: {ws.title} (max_row={ws.max_row}, max_col={ws.max_column})")
        count = 0
        for row in ws.iter_rows(values_only=True):
            vals = [clean(str(c)) for c in row if c is not None]
            if vals:
                parts.append(" | ".join(vals)[:300])
                count += 1
            if count >= max_rows:
                break
        # 注意: 很多调研表前几行只有表头, 数据在更后; 这里只取结构
    return "\n".join(parts)


def extract_docx(path, max_paras=60, max_tables=3, max_rows_per_table=5):
    doc = Document(path)
    parts = []
    for p in doc.paragraphs:
        t = clean(p.text)
        if not t:
            continue
        style = p.style.name if p.style else ""
        prefix = "## " if "Heading 1" in style else ("### " if "Heading" in style else "")
        parts.append(prefix + t)
        if len(parts) >= max_paras:
            break
    for ti, tbl in enumerate(doc.tables):
        parts.append(f"\n[Table {ti+1}]")
        for ri, row in enumerate(tbl.rows):
            cells = [clean(c.text) for c in row.cells]
            parts.append(" | ".join(cells)[:300])
            if ri >= max_rows_per_table:
                break
        if ti >= max_tables - 1:
            break
    return "\n".join(parts)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("dir", help="文档目录")
    ap.add_argument("--out", default="_INDEX.md")
    args = ap.parse_args()

    files = sorted(glob.glob(os.path.join(args.dir, "*")))
    files = [f for f in files if not os.path.basename(f) == args.out]
    index = [f"# 文档资料索引\n", f"文件总数: {len(files)}\n"]
    for f in files:
        name = os.path.basename(f)
        if name == args.out:
            continue
        index.append(f"\n---\n## 📄 {name}\n")
        try:
            if name.endswith(".xlsx"):
                index.append(extract_xlsx(f))
            elif name.endswith(".docx"):
                index.append(extract_docx(f))
            else:
                index.append("(非 xlsx/docx, 跳过)")
        except Exception as e:
            index.append(f"[提取失败: {e}]")
    out_path = os.path.join(args.dir, args.out)
    with open(out_path, "w", encoding="utf-8") as fh:
        fh.write("\n".join(index))
    print(f"[ok] 索引已生成: {out_path} ({os.path.getsize(out_path)} 字节, {len(files)} 文件)")


if __name__ == "__main__":
    main()

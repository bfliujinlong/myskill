#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
把模型1 的大数据评估单元回写进《应用复杂度评估 V7》体系工作簿的「应用评估」页。

只填 A-G（ID/名称/序号/D1-D4）、M-Q（五个修正系数）、AC（备注）、AD（团队），
不动原表任何公式，让 H-L、R-AB 由原生公式链自行计算。

用法:
  python3 inject_v7.py --workbook "应用复杂度评估 V7.xlsx" --config input.json \
                       --out "应用复杂度评估 V7-含大数据评估.xlsx" --team "团队4-大数据"
"""
import argparse, json, shutil, sys, os

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

THIN = Side(style='thin', color='BFBFBF')
BD = Border(THIN, THIN, THIN, THIN)
FILL = PatternFill('solid', fgColor='FFF2CC')

# 应用评估页列位
C_ID, C_NAME, C_SEQ = 1, 2, 3
C_D1, C_D4 = 4, 7
C_AUTO, C_DOC, C_BW, C_COMP, C_OTHER = 13, 14, 15, 16, 17
C_NOTE, C_TEAM = 29, 30


def die(m):
    sys.stderr.write("BLOCKED: %s\n" % m); sys.exit(2)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--workbook', required=True)
    ap.add_argument('--config', required=True)
    ap.add_argument('--out', required=True)
    ap.add_argument('--team', default='团队4-大数据')
    ap.add_argument('--sheet', default='应用评估')
    ap.add_argument('--start-row', type=int, default=0, help='0=自动接在最后一行数据之后')
    a = ap.parse_args()

    for p in (a.workbook, a.config):
        if not os.path.exists(p):
            die("文件不存在: %s" % p)
    with open(a.config, encoding='utf-8') as f:
        cfg = json.load(f)
    units = cfg.get('model1_units') or []
    if not units:
        die("config 里没有 model1_units")
    co = cfg.get('coefficients') or {}

    shutil.copy(a.workbook, a.out)
    wb = openpyxl.load_workbook(a.out)
    if a.sheet not in wb.sheetnames:
        die("工作簿里没有 '%s' 页" % a.sheet)
    ws = wb[a.sheet]

    # 定位最后一行有应用名称的数据行
    last, last_seq = 1, 0
    for r in range(2, ws.max_row + 1):
        if ws.cell(r, C_NAME).value not in (None, ''):
            last = r
            s = ws.cell(r, C_SEQ).value
            if isinstance(s, str) and s.startswith('应用'):
                try: last_seq = max(last_seq, int(s[2:]))
                except ValueError: pass
    start = a.start_row or (last + 1)

    for i, u in enumerate(units):
        r = start + i
        ov = u.get('coef_override', {})
        def cf(k, default_from_cfg=True):
            if k in ov: return float(ov[k])
            return float(co.get(k, 1.0)) if default_from_cfg else 1.0
        ws.cell(r, C_ID, u.get('id', 'BD-%02d' % (i + 1)))
        ws.cell(r, C_NAME, (('%s-%s' % (u['group'], u['name'])) if u.get('group') else u['name']))
        ws.cell(r, C_SEQ, '应用%d' % (last_seq + i + 1))
        for j, k in enumerate(('d1', 'd2', 'd3', 'd4')):
            ws.cell(r, C_D1 + j, u[k])
        ws.cell(r, C_AUTO, cf('automation'))
        ws.cell(r, C_DOC, cf('doc'))
        # 带宽系数默认计在模型2，本表取 1.0，除非显式要求
        ws.cell(r, C_BW, cf('bandwidth') if cfg.get('model1_include_bandwidth') else 1)
        ws.cell(r, C_COMP, cf('compliance'))
        ws.cell(r, C_OTHER, cf('other'))
        ws.cell(r, C_NOTE, u.get('note', ''))
        ws.cell(r, C_TEAM, a.team)
        for c in range(1, C_TEAM + 1):
            cc = ws.cell(r, c)
            cc.border = BD
            cc.alignment = Alignment(vertical='center', wrap_text=(c == C_NOTE))
            if c in (C_ID, C_NAME, C_SEQ, C_D1, C_D1 + 1, C_D1 + 2, C_D4,
                     C_AUTO, C_DOC, C_BW, C_COMP, C_OTHER, C_NOTE, C_TEAM):
                cc.fill = FILL

    # 清掉模板残留的 D-G 占位值，避免空行被当成 L1 参与汇总
    for r in range(start + len(units), ws.max_row + 1):
        if ws.cell(r, C_NAME).value in (None, ''):
            for c in range(C_D1, C_D4 + 1):
                ws.cell(r, c).value = None

    # 团队汇总页：把新团队名挂到第一个空的团队行，便于原表 SUMIF 汇总
    if '团队汇总' in wb.sheetnames:
        t = wb['团队汇总']
        for rr in range(5, 106):
            if t.cell(rr, 1).value in (None, ''):
                t.cell(rr, 1, a.team)
                break

    # 说明页
    title = '大数据评估说明'
    if title in wb.sheetnames:
        del wb[title]
    s = wb.create_sheet(title, wb.sheetnames.index(a.sheet) + 1)
    doc = cfg.get('coefficients', {}).get('doc', 1.0)
    rows = [
        ['%s — 大数据评估口径说明（我方口径）' % cfg.get('project', ''), ''],
        ['', ''],
        ['一、与“%s”页的关系' % a.sheet, ''],
        ['', '第 %d-%d 行为大数据评估单元（%s），共 %d 个，沿用本表 D1-D4 → 等级 → 基准人天 → 修正系数 的既有模型。'
             % (start, start + len(units) - 1, a.team, len(units))],
        ['', '大数据总人力 = 模型1（本表算出）+ 模型2（标准服务目录）+ 模型3（超服务目录项），后两层见独立评估工作簿。'],
        ['', ''],
        ['二、D1-D4 在大数据场景的定义', ''],
        ['架构复杂度 D1', '源端托管服务→目标端的异构程度：同构搬迁1；组件替换3；数仓+计算引擎+编排整栈重构5'],
        ['数据复杂度 D2', '数据量级、表数、元数据规模、Schema/SQL方言差异：单库小表1；万表级元数据+方言重写5'],
        ['耦合度 D3', '被多少上下游依赖、是否公共SDK或调度收口：独立应用1；全平台依赖的SDK/调度器5'],
        ['改造深度 D4', 'Rehost 1；Replatform 3；代码级 Refactor（SQL/UDF/作业逻辑重写）5'],
        ['', ''],
        ['三、本次修正系数取值', ''],
        ['自动化系数', str(co.get('automation', 1.0)) + ' — 异构改造无可用全自动工具时取 1.0'],
        ['文档系数', str(doc) + ' — 1.5 表示客户未开放源码仓库/架构文档；开放后可降至 1.0'],
        ['带宽系数', '1.0 — 数据搬迁的带宽影响单列在模型2 测算，本表不重复计入'],
        ['合规系数', str(co.get('compliance', 1.0))],
        ['', ''],
        ['四、判重规则', ''],
        ['', '本表按模块整体计价的应用，其内部作业不再进入模型3；模型3 只收未被本表覆盖的作业/对象级工作。'],
    ]
    for i, row in enumerate(rows, 1):
        for j, v in enumerate(row, 1):
            s.cell(i, j, v)
    s.column_dimensions['A'].width = 18
    s.column_dimensions['B'].width = 110
    s['A1'].font = Font(bold=True, size=13)
    for i in (3, 7, 13, 19):
        s.cell(i, 1).font = Font(bold=True, color='C00000')
    for i in range(1, len(rows) + 1):
        s.cell(i, 2).alignment = Alignment(wrap_text=True, vertical='center')
    s.sheet_properties.tabColor = 'C00000'

    wb.save(a.out)
    print("写入 %d 个大数据评估单元，行 %d-%d" % (len(units), start, start + len(units) - 1))
    print("CREATED: %s" % a.out)
    print("提示：openpyxl 不计算公式，请用 Excel/WPS 打开一次触发重算，或用 LibreOffice --convert-to xlsx 重算。")


if __name__ == '__main__':
    main()

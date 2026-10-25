#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
大数据迁移评估 — 确定性三层模型计算器

用法:
    python3 bd_assess.py --config input.json --out 输出.xlsx
    python3 bd_assess.py --config input.json --print-only

输入契约见 examples/envision-input.json。
纪律: 缺必填项 -> 退出码 2, 拒绝产出。不猜数、不补默认单价。
"""
import argparse, json, sys, os

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.stderr.write("BLOCKED: 需要 openpyxl (pip install openpyxl)\n")
    sys.exit(3)

# ---------------- 基准常量 (与 references/baseline-params.md 同源) ----------------
MD_PER_MM = 22.0

LEVELS = [  # (名称, 下限含, 上限不含, 低人天, 高人天)
    ('L1-简单型',   1.0, 1.8,  5,  7),
    ('L2-普通型',   1.8, 2.6,  7, 12),
    ('L3-中等型',   2.6, 3.4, 15, 25),
    ('L4-复杂型',   3.4, 4.2, 30, 50),
    ('L5-极高难度', 4.2, 6.0, 60, 90),
]

PHASE_SETS = {
    'v7': [('调研与方案设计', 0.20), ('环境准备与迁移实施', 0.30),
           ('应用适配与调试', 0.30), ('验收与割接', 0.20)],
    'six': [('调研', 0.10), ('方案设计及评审', 0.30), ('方案验证及复盘', 0.15),
            ('迁移实施', 0.30), ('业务割接', 0.10), ('迁移验收', 0.05)],
    'bigdata': [('调研', 0.110), ('方案设计及评审', 0.170), ('方案验证及复盘', 0.110),
                ('迁移实施', 0.500), ('业务割接', 0.055), ('迁移验收', 0.055)],
}

ROLES = [  # (岗位, 比例, 成本价, 刊例价)
    ('项目经理 PM',            0.1329, 3300, 4200),
    ('云架构师 Architecture',  0.2601, 3300, 4200),
    ('云工程师 Cloud Engineer', 0.5256, 1900, 2500),
    ('DBA',                    0.0814, 2500, 3300),
]

# 模型2 标准服务目录：item -> (产能, 单位, 计费量上限 或 None)
CATALOG = {
    '主机实例':            (40,   '台/人天',  None),
    '容器POD':             (15,   '个/人天',  None),
    '镜像仓库':            (10,   '个/人天',  None),
    '对象存储(按桶)':      (10,   '个/人天',  None),
    '对象存储(按量)':      (2000, 'GB/人天',  102400),
    '存储迁移(非对象)':    (1000, 'GB/人天',  102400),
    '中间件(按套)':        (2,    '套/人天',  None),
    '中间件(按量)':        (500,  'GB/人天',  5120),
    '数据库(按套)':        (2,    '套/人天',  None),
    '数据库(按量)':        (800,  'GB/人天',  51200),
}

# ---------------- 样式 ----------------
THIN = Side(style='thin', color='BFBFBF')
BD = Border(THIN, THIN, THIN, THIN)
H_FILL = PatternFill('solid', fgColor='305496')
H_FONT = Font(bold=True, color='FFFFFF', size=10)
WARN_FILL = PatternFill('solid', fgColor='FFC7CE')
OK_FILL = PatternFill('solid', fgColor='E2EFDA')
T_FILL = PatternFill('solid', fgColor='FFF2CC')


def die(msg):
    sys.stderr.write("BLOCKED: %s\n" % msg)
    sys.exit(2)


def r1(x):
    return round(x + 1e-9, 2)


# ---------------- 模型1 ----------------
def level_of(avg):
    for name, lo, hi, dl, dh in LEVELS:
        if lo <= avg < hi:
            return name, dl, dh
    return LEVELS[-1][0], LEVELS[-1][3], LEVELS[-1][4]


def calc_model1(cfg):
    units = cfg.get('model1_units') or []
    co = cfg.get('coefficients') or {}
    # 模型1 默认不计带宽系数（计在模型2），除非显式 model1_include_bandwidth
    keys = ['automation', 'doc', 'compliance', 'other']
    if cfg.get('model1_include_bandwidth'):
        keys.insert(2, 'bandwidth')
    rows, tot_lo, tot_hi = [], 0.0, 0.0
    for u in units:
        for k in ('name', 'd1', 'd2', 'd3', 'd4'):
            if u.get(k) in (None, ''):
                die("模型1 单元缺字段 %s: %s" % (k, u))
        ds = [u['d1'], u['d2'], u['d3'], u['d4']]
        for d in ds:
            if not (1 <= d <= 5):
                die("D 分必须 1-5: %s" % u)
        avg = sum(ds) / 4.0
        lv, bl, bh = level_of(avg)
        coef = 1.0
        for k in keys:
            coef *= float(u.get('coef_override', {}).get(k, co.get(k, 1.0)))
        lo, hi = r1(bl * coef), r1(bh * coef)
        tot_lo += lo
        tot_hi += hi
        rows.append(dict(group=u.get('group', ''), name=u['name'], d1=ds[0], d2=ds[1],
                         d3=ds[2], d4=ds[3], total=sum(ds), avg=r1(avg), level=lv,
                         base_lo=bl, base_hi=bh, coef=r1(coef), lo=lo, hi=hi,
                         note=u.get('note', '')))
    return rows, r1(tot_lo), r1(tot_hi)


# ---------------- 模型2 ----------------
def calc_model2(cfg):
    items = cfg.get('model2_resources') or []
    co = cfg.get('coefficients') or {}
    bw = float(co.get('bandwidth', 1.0))
    inside, over = [], []
    for it in items:
        name = it.get('item')
        qty = it.get('qty')
        if name is None or qty is None:
            die("模型2 条目缺 item/qty: %s" % it)
        if 'rate' in it:
            rate, unit, cap = it['rate'], it.get('rate_unit', ''), it.get('cap')
        elif name in CATALOG:
            rate, unit, cap = CATALOG[name]
        else:
            die("模型2 条目 '%s' 不在标准服务目录，且未显式给 rate。"
                "若确属目录外工作，请移到 model3_items。" % name)
        md = qty / float(rate)
        rec = dict(item=name, qty=qty, unit=it.get('unit', ''), rate=rate,
                   rate_unit=unit, cap=cap, md=r1(md), note=it.get('note', ''))
        if cap and qty > cap:
            rec['over_x'] = r1(qty / float(cap))
            over.append(rec)
        else:
            inside.append(rec)
    md_inside = r1(sum(x['md'] for x in inside))
    md_inside_adj = r1(md_inside * bw)

    # 超上限项替代算法
    alt = cfg.get('model2_alternatives') or []
    md_alt = 0.0
    alt_rows = []
    for a in alt:
        v = a.get('man_days')
        if v is None:
            die("model2_alternatives 条目缺 man_days: %s" % a)
        md_alt += float(v)
        alt_rows.append(dict(item=a.get('item', ''), md=r1(float(v)),
                             note=a.get('basis', '')))
    md_alt = r1(md_alt)
    return inside, over, alt_rows, md_inside, bw, md_inside_adj, md_alt, r1(md_inside_adj + md_alt)


def bandwidth_table(total_gb, scenarios):
    out = []
    for s in scenarios:
        gbps = float(s['gbps']); eff = float(s.get('eff', 0.7)); hrs = float(s.get('hours', 18))
        per_day_gb = gbps / 8.0 * 3600.0 * hrs * eff
        days = total_gb / per_day_gb if per_day_gb else 0
        out.append(dict(name=s.get('name', ''), gbps=gbps, eff=eff, hours=hrs,
                        per_day_tb=r1(per_day_gb / 1024.0), days=r1(days),
                        months=r1(days / 30.0)))
    return out


# ---------------- 模型3 ----------------
def calc_model3(cfg):
    rows, tot = [], 0.0
    for it in cfg.get('model3_items') or []:
        if 'man_days' in it:
            md = float(it['man_days'])
            qty, rate, ru = it.get('qty', ''), it.get('rate', ''), it.get('rate_unit', '')
        else:
            for k in ('qty', 'rate'):
                if it.get(k) in (None, ''):
                    die("模型3 条目缺 %s 且未给 man_days: %s" % (k, it))
            qty, rate, ru = it['qty'], it['rate'], it.get('rate_unit', '')
            md = qty / float(rate) + float(it.get('fixed_md', 0))
        tot += md
        rows.append(dict(cat=it.get('category', ''), item=it.get('item', ''),
                         qty=qty, unit=it.get('unit', ''), rate=rate, rate_unit=ru,
                         md=r1(md), basis=it.get('basis', '')))
    return rows, r1(tot)


# ---------------- 汇总 ----------------
def summarize(m1_lo, m1_hi, m2, m3, phase_key, price_mode):
    lo = r1(m1_lo + m2 + m3)
    hi = r1(m1_hi + m2 + m3)
    mid = r1((lo + hi) / 2.0)
    phases = [(n, w, r1(mid * w)) for n, w in PHASE_SETS[phase_key]]
    roles = []
    for name, pct, cost, list_p in ROLES:
        md = r1(mid * pct)
        roles.append(dict(name=name, pct=pct, md=md,
                          cost=round(md * cost), listp=round(md * list_p),
                          cost_rate=cost, list_rate=list_p))
    total_cost = sum(r['cost'] for r in roles)
    total_list = sum(r['listp'] for r in roles)
    return dict(lo=lo, mid=mid, hi=hi, mm=r1(mid / MD_PER_MM),
                phases=phases, roles=roles,
                total_cost=total_cost, total_list=total_list)


# ---------------- 写表 ----------------
def hdr(ws, row, cols, widths=None):
    for i, c in enumerate(cols, 1):
        cell = ws.cell(row, i, c)
        cell.fill = H_FILL; cell.font = H_FONT; cell.border = BD
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    if widths:
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w


def put(ws, row, vals, fill=None, bold=False, wrap_cols=()):
    for i, v in enumerate(vals, 1):
        c = ws.cell(row, i, v)
        c.border = BD
        c.alignment = Alignment(vertical='center', wrap_text=(i in wrap_cols))
        if fill: c.fill = fill
        if bold: c.font = Font(bold=True)


def build(cfg, out, m1, m1lo, m1hi, m2in, m2over, m2alt, m2_raw, bw, m2_adj, m2_altmd, m2, m3, m3md, S):
    wb = openpyxl.Workbook()
    proj = cfg.get('project', '未命名项目')

    # 01 说明
    ws = wb.active; ws.title = '01-评估模型说明'
    ws.column_dimensions['A'].width = 22; ws.column_dimensions['B'].width = 110
    lines = [
        ('%s — 大数据迁移人力评估（我方口径）' % proj, ''),
        ('', ''),
        ('口径声明', '本表完全采用我方三层评估模型独立测算，未沿用客户/友商的人月数字。对方数字仅在“差异对比”中作为参照。'),
        ('', ''),
        ('模型1 复杂度分级', 'D1架构/D2数据/D3耦合/D4改造深度 各1-5分 → 平均分落等级 → 基准人天区间 × 修正系数(连乘)。覆盖平台/模块级改造。'),
        ('模型2 标准服务目录', '资源量 ÷ 标准产能 = 人天。超出目录计费上限的量走替代算法（工具人天 + 传输期值守），不做线性外推。'),
        ('模型3 超服务目录项', '标准服务目录中查不到计费项的工作，按产能定律逐项定价，全部需单独商务谈判。覆盖作业/对象级改造。'),
        ('', ''),
        ('分层判重', '模型1 按模块整体计价的应用，其内部作业不再进入模型3；模型3 只收未被模型1 覆盖的作业/对象级工作。'),
        ('带宽系数', '只计一次，默认计在模型2；模型1 取 1.0。'),
        ('换算', '1 人月 = 22 人天；中值 = (低+高)/2。'),
        ('阶段口径', cfg.get('phase_set', 'v7')),
        ('价格口径', '成本价与刊例价并列输出，取用时必须声明取的是哪一档。'),
        ('', ''),
        ('本次结果', '总计 低 %.2f / 中 %.2f / 高 %.2f 人天 = %.2f 人月；成本价 ¥%s；刊例价 ¥%s'
         % (S['lo'], S['mid'], S['hi'], S['mm'], format(S['total_cost'], ','), format(S['total_list'], ','))),
        ('  模型1', '低 %.2f / 中 %.2f / 高 %.2f 人天' % (m1lo, (m1lo + m1hi) / 2, m1hi)),
        ('  模型2', '目录内 %.2f × 带宽系数 %.1f = %.2f，+ 超上限替代 %.2f = %.2f 人天' % (m2_raw, bw, m2_adj, m2_altmd, m2)),
        ('  模型3', '%.2f 人天' % m3md),
    ]
    for i, (a, b) in enumerate(lines, 1):
        ws.cell(i, 1, a).font = Font(bold=True)
        ws.cell(i, 2, b).alignment = Alignment(wrap_text=True, vertical='center')
    ws['A1'].font = Font(bold=True, size=14)
    ws.sheet_properties.tabColor = '305496'

    # 02 模型1
    ws = wb.create_sheet('02-模型1·复杂度分级')
    hdr(ws, 1, ['平台/分组', '模块', 'D1架构', 'D2数据', 'D3耦合', 'D4深度', '总分', '平均分', '等级',
                '基准低', '基准高', '修正系数', '调整后低', '调整后高', '打分依据'],
        [14, 26, 7, 7, 7, 7, 7, 8, 13, 8, 8, 9, 10, 10, 60])
    r = 2
    for x in m1:
        put(ws, r, [x['group'], x['name'], x['d1'], x['d2'], x['d3'], x['d4'], x['total'],
                    x['avg'], x['level'], x['base_lo'], x['base_hi'], x['coef'],
                    x['lo'], x['hi'], x['note']], wrap_cols=(15,))
        r += 1
    put(ws, r, ['合计', '%d 个单元' % len(m1), '', '', '', '', '', '', '', '', '', '',
                m1lo, m1hi, '中值 %.2f 人天' % ((m1lo + m1hi) / 2)], fill=T_FILL, bold=True)
    ws.freeze_panes = 'C2'; ws.sheet_properties.tabColor = '70AD47'

    # 03 模型2
    ws = wb.create_sheet('03-模型2·标准服务目录')
    hdr(ws, 1, ['计费项', '数量', '单位', '标准产能', '产能单位', '计费上限', '人天', '状态', '备注'],
        [24, 14, 8, 11, 12, 12, 10, 16, 52])
    r = 2
    ws.cell(r, 1, '一、目录内项').font = Font(bold=True, color='C00000'); r += 1
    for x in m2in:
        put(ws, r, [x['item'], x['qty'], x['unit'], x['rate'], x['rate_unit'],
                    x['cap'] or '—', x['md'], '目录内', x['note']], fill=OK_FILL, wrap_cols=(9,))
        r += 1
    put(ws, r, ['目录内小计', '', '', '', '', '', m2_raw, '', ''], fill=T_FILL, bold=True); r += 1
    put(ws, r, ['× 带宽系数 %.1f' % bw, '', '', '', '', '', m2_adj, '', '搬迁带宽影响计入本层，模型1 不重复计'],
        fill=T_FILL, bold=True); r += 2

    ws.cell(r, 1, '二、超计费上限项（不得线性外推，走替代算法）').font = Font(bold=True, color='C00000'); r += 1
    for x in m2over:
        put(ws, r, [x['item'], x['qty'], x['unit'], x['rate'], x['rate_unit'], x['cap'],
                    x['md'], '超上限 %.1f×' % x['over_x'],
                    '线性外推结果 %.2f 人天不成立，改用替代算法' % x['md']], fill=WARN_FILL, wrap_cols=(9,))
        r += 1
    r += 1
    ws.cell(r, 1, '三、替代算法（工具操作 + 传输期值守）').font = Font(bold=True, color='C00000'); r += 1
    for x in m2alt:
        put(ws, r, [x['item'], '', '', '', '', '', x['md'], '替代算法', x['note']], wrap_cols=(9,)); r += 1
    put(ws, r, ['替代算法小计', '', '', '', '', '', m2_altmd, '', ''], fill=T_FILL, bold=True); r += 1
    put(ws, r, ['模型2 合计', '', '', '', '', '', m2, '', '= 目录内(含带宽系数) + 替代算法'],
        fill=T_FILL, bold=True); r += 2

    bws = cfg.get('bandwidth_scenarios')
    if bws and cfg.get('transfer_total_gb'):
        ws.cell(r, 1, '四、搬迁日历工期场景（工期风险常大于人力风险）').font = Font(bold=True, color='C00000'); r += 1
        hdr(ws, r, ['场景', '带宽Gbps', '有效率', '每日小时', '日吞吐TB', '日历天', '折合月', '', '']); r += 1
        for s in bandwidth_table(float(cfg['transfer_total_gb']), bws):
            put(ws, r, [s['name'], s['gbps'], s['eff'], s['hours'], s['per_day_tb'],
                        s['days'], s['months'], '', '']); r += 1
    ws.sheet_properties.tabColor = 'FFC000'

    # 04 模型3
    ws = wb.create_sheet('04-模型3·超服务目录项')
    hdr(ws, 1, ['类别', '工作项', '数量', '单位', '标准产能', '产能单位', '人天', '测算依据'],
        [16, 32, 12, 8, 11, 12, 10, 58])
    r = 2
    for x in m3:
        put(ws, r, [x['cat'], x['item'], x['qty'], x['unit'], x['rate'], x['rate_unit'],
                    x['md'], x['basis']], wrap_cols=(8,)); r += 1
    put(ws, r, ['合计', '%d 项' % len(m3), '', '', '', '', m3md,
                '本层全部无标准计费项，需单独商务谈判'], fill=WARN_FILL, bold=True)
    ws.freeze_panes = 'C2'; ws.sheet_properties.tabColor = 'C00000'

    # 05 汇总
    ws = wb.create_sheet('05-汇总·阶段·岗位·成本')
    ws.column_dimensions['A'].width = 26
    for col in 'BCDEFG': ws.column_dimensions[col].width = 16
    ws.column_dimensions['H'].width = 46
    r = 1
    ws.cell(r, 1, '一、三层汇总').font = Font(bold=True, color='C00000'); r += 1
    hdr(ws, r, ['层级', '低(人天)', '中(人天)', '高(人天)', '人月(中)', '占比', '', '说明']); r += 1
    m1mid = r1((m1lo + m1hi) / 2)
    for nm, a, b, c, desc in [
        ('模型1 复杂度分级', m1lo, m1mid, m1hi, '平台/模块级改造'),
        ('模型2 标准服务目录', m2, m2, m2, '资源搬迁，含带宽系数与超上限替代算法'),
        ('模型3 超服务目录项', m3md, m3md, m3md, '无标准计费项，需单独商务谈判'),
    ]:
        put(ws, r, [nm, a, b, c, r1(b / MD_PER_MM), r1(b / S['mid'] * 100), '', desc]); r += 1
    put(ws, r, ['合计', S['lo'], S['mid'], S['hi'], S['mm'], 100.0, '', ''], fill=T_FILL, bold=True); r += 2

    ws.cell(r, 1, '二、阶段拆分（按中值 %.2f 人天，口径 %s）' % (S['mid'], cfg.get('phase_set', 'v7'))).font = Font(bold=True, color='C00000'); r += 1
    hdr(ws, r, ['阶段', '权重', '人天', '人月', '', '', '', '']); r += 1
    for n, w, md in S['phases']:
        put(ws, r, [n, w, md, r1(md / MD_PER_MM), '', '', '', '']); r += 1
    r += 1

    ws.cell(r, 1, '三、岗位拆分与成本').font = Font(bold=True, color='C00000'); r += 1
    hdr(ws, r, ['岗位', '比例', '人天', '成本单价', '成本金额', '刊例单价', '刊例金额', '']); r += 1
    for x in S['roles']:
        put(ws, r, [x['name'], x['pct'], x['md'], x['cost_rate'], x['cost'],
                    x['list_rate'], x['listp'], '']); r += 1
    put(ws, r, ['合计', 1.0, S['mid'], '', S['total_cost'], '', S['total_list'], ''],
        fill=T_FILL, bold=True); r += 2

    cmp_rows = cfg.get('comparisons') or []
    if cmp_rows:
        ws.cell(r, 1, '四、差异对比与根因').font = Font(bold=True, color='C00000'); r += 1
        hdr(ws, r, ['口径', '人天', '人月', '金额', '与我方差异(人天)', '倍数', '', '根因']); r += 1
        for c in cmp_rows:
            md = c.get('man_days')
            delta = r1(S['mid'] - md) if md else ''
            mult = r1(S['mid'] / md) if md else ''
            put(ws, r, [c.get('name', ''), md, r1(md / MD_PER_MM) if md else '',
                        c.get('amount', ''), delta, mult, '', c.get('root_cause', '')],
                fill=WARN_FILL, wrap_cols=(8,)); r += 1
        r += 1

    lev = cfg.get('levers') or []
    if lev:
        ws.cell(r, 1, '五、降本杠杆').font = Font(bold=True, color='C00000'); r += 1
        hdr(ws, r, ['杠杆', '客户动作', '节省人天', '节省人月', '节省金额(成本价)', '', '', '']); r += 1
        for l in lev:
            md = float(l.get('save_md', 0))
            put(ws, r, [l.get('name', ''), l.get('action', ''), r1(md), r1(md / MD_PER_MM),
                        l.get('save_amount', ''), '', '', '']); r += 1
    ws.sheet_properties.tabColor = '7030A0'

    # 06 清单
    ws = wb.create_sheet('06-待客户确认清单')
    hdr(ws, 1, ['优先级', '需要的信息', '缺失后果', '状态'], [10, 48, 66, 14])
    r = 2
    for it in cfg.get('open_items') or []:
        p = it.get('priority', 'P1')
        put(ws, r, [p, it.get('item', ''), it.get('impact', ''), it.get('status', '待客户提供')],
            fill=WARN_FILL if p == 'P0' else None, wrap_cols=(2, 3)); r += 1
    ws.freeze_panes = 'A2'; ws.sheet_properties.tabColor = 'ED7D31'

    wb.save(out)
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--config', required=True)
    ap.add_argument('--out')
    ap.add_argument('--print-only', action='store_true')
    a = ap.parse_args()

    if not os.path.exists(a.config):
        die("配置文件不存在: %s" % a.config)
    with open(a.config, encoding='utf-8') as f:
        cfg = json.load(f)

    ps = cfg.get('phase_set', 'v7')
    if ps not in PHASE_SETS:
        die("phase_set 必须是 %s" % list(PHASE_SETS))
    pm = cfg.get('price_mode', 'cost')
    if not (cfg.get('model1_units') or cfg.get('model2_resources') or cfg.get('model3_items')):
        die("三层输入全空，无可测算内容")

    m1, m1lo, m1hi = calc_model1(cfg)
    m2in, m2over, m2alt, m2raw, bw, m2adj, m2altmd, m2 = calc_model2(cfg)
    m3, m3md = calc_model3(cfg)
    S = summarize(m1lo, m1hi, m2, m3md, ps, pm)

    print("模型1  低 %.2f  中 %.2f  高 %.2f" % (m1lo, (m1lo + m1hi) / 2, m1hi))
    print("模型2  目录内 %.2f ×%.1f = %.2f  + 替代 %.2f = %.2f" % (m2raw, bw, m2adj, m2altmd, m2))
    if m2over:
        for x in m2over:
            print("  !! 超计费上限: %s 量 %s 超 %.1f×，线性外推 %.2f 人天不成立" %
                  (x['item'], x['qty'], x['over_x'], x['md']))
    print("模型3  %.2f （%d 项，全部需单独商务谈判）" % (m3md, len(m3)))
    print("合计   低 %.2f  中 %.2f  高 %.2f 人天 = %.2f 人月" % (S['lo'], S['mid'], S['hi'], S['mm']))
    print("成本价 ¥%s   刊例价 ¥%s" % (format(S['total_cost'], ','), format(S['total_list'], ',')))

    if a.print_only:
        return
    if not a.out:
        die("缺 --out")
    p = build(cfg, a.out, m1, m1lo, m1hi, m2in, m2over, m2alt, m2raw, bw, m2adj, m2altmd, m2, m3, m3md, S)
    print("CREATED: %s" % p)


if __name__ == '__main__':
    main()

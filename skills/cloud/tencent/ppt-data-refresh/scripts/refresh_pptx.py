#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
refresh_pptx.py — 刷新 PowerPoint 中的业务数据，保留原样式。

用法:
    python refresh_pptx.py --deck Q3.pptx --scan --out inventory.json
    python refresh_pptx.py --deck Q3.pptx --data data.json --dry-run --report diff.md
    python refresh_pptx.py --deck Q3.pptx --data data.json --out Q4.pptx --report diff.md

支持:
    - 文本框 / 占位符中的 {{TOKEN}} 替换
    - 表格单元格中的 {{TOKEN}} 替换
    - 位置映射替换（--map）
    - 图表 series 数据更新（data.json 的 charts 段）
    - 疑似数据点扫描（千分位数字、百分比、金额）
    - 变更报告输出

依赖: pip install python-pptx
"""

import argparse
import json
import re
import sys
from pathlib import Path

try:
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
except ImportError:
    print(json.dumps({
        "ok": False,
        "error": "缺少依赖 python-pptx，请先执行: pip install python-pptx"
    }, ensure_ascii=False, indent=2))
    sys.exit(2)

TOKEN_RE = re.compile(r"\{\{\s*([A-Za-z0-9_.\-]+)\s*\}\}")

# 疑似数据点：千分位数字、百分比、金额、带小数的量
DATAISH_RE = re.compile(
    r"(?:[$¥￥€]\s?\d[\d,]*(?:\.\d+)?\s?(?:万|亿|M|B|K|k)?"
    r"|\d[\d,]{2,}(?:\.\d+)?\s?(?:万|亿|M|B|K|k)?"
    r"|\d+(?:\.\d+)?\s?%)"
)


def iter_text_blocks(prs):
    """产出 (slide_no, shape_name, kind, text_frame)。kind: text / table / notes / chart"""
    for idx, slide in enumerate(prs.slides, start=1):
        for shape in slide.shapes:
            name = getattr(shape, "name", "unnamed")
            if shape.has_text_frame:
                yield idx, name, "text", shape.text_frame
            if getattr(shape, "has_table", False) and shape.has_table:
                for r_i, row in enumerate(shape.table.rows):
                    for c_i, cell in enumerate(row.cells):
                        yield idx, "%s[r%d,c%d]" % (name, r_i, c_i), "table", cell.text_frame
            if getattr(shape, "has_chart", False) and shape.has_chart:
                yield idx, name, "chart", shape.chart
        if slide.has_notes_slide:
            notes_tf = slide.notes_slide.notes_text_frame
            if notes_tf is not None:
                yield idx, "notes", "notes", notes_tf


def replace_in_paragraph(paragraph, tokens, changes, locator):
    """在段落内替换 {{TOKEN}}，尽量保留 run 级样式。"""
    runs = paragraph.runs
    if not runs:
        return
    full = "".join(r.text for r in runs)
    if "{{" not in full:
        return

    def _sub(match):
        key = match.group(1)
        if key not in tokens:
            return match.group(0)
        changes.append({
            "slide": locator[0],
            "shape": locator[1],
            "old": match.group(0),
            "new": str(tokens[key]),
            "unmatched": False,
        })
        return str(tokens[key])

    new_full = TOKEN_RE.sub(_sub, full)
    if new_full == full:
        return

    # 判断所有 token 是否都落在单个 run 内（可安全保留样式）
    single_run_safe = True
    for match in TOKEN_RE.finditer(full):
        start, end = match.span()
        pos, ok = 0, False
        for run in runs:
            run_end = pos + len(run.text)
            if start >= pos and end <= run_end:
                ok = True
                break
            pos = run_end
        if not ok:
            single_run_safe = False
            break

    if single_run_safe:
        for run in runs:
            if "{{" in run.text:
                run.text = TOKEN_RE.sub(
                    lambda m: str(tokens.get(m.group(1), m.group(0))), run.text
                )
    else:
        # 跨 run 的 token：整段写入首个 run，其余清空
        runs[0].text = new_full
        for run in runs[1:]:
            run.text = ""


def scan_deck(prs):
    inventory = {"tokens": [], "dataish": [], "charts": []}
    seen_tokens = set()
    for slide_no, name, kind, obj in iter_text_blocks(prs):
        if kind == "chart":
            try:
                chart = obj
                series_names = [s.name for s in chart.series]
                cats = [str(c) for c in chart.plots[0].categories]
                inventory["charts"].append({
                    "id": "slide%d/%s" % (slide_no, name),
                    "series": series_names,
                    "categories": cats,
                })
            except Exception as exc:  # noqa: BLE001
                inventory["charts"].append({
                    "id": "slide%d/%s" % (slide_no, name),
                    "error": str(exc),
                })
            continue

        text = obj.text if kind == "notes" else "\n".join(p.text for p in obj.paragraphs)
        for match in TOKEN_RE.finditer(text):
            key = match.group(1)
            if key not in seen_tokens:
                seen_tokens.add(key)
                inventory["tokens"].append({"token": key, "first_seen": "slide%d/%s" % (slide_no, name)})
        for match in DATAISH_RE.finditer(text):
            inventory["dataish"].append({
                "slide": slide_no,
                "shape": name,
                "value": match.group(0).strip(),
            })
    return inventory


def apply_tokens(prs, tokens, changes):
    for slide_no, name, kind, obj in iter_text_blocks(prs):
        if kind == "chart":
            continue
        if kind == "notes":
            paragraphs = obj.paragraphs
        else:
            paragraphs = obj.paragraphs
        for paragraph in paragraphs:
            replace_in_paragraph(paragraph, tokens, changes, (slide_no, name))


def apply_mapping(prs, mapping, changes):
    by_slide = {}
    for item in mapping:
        by_slide.setdefault(int(item["slide"]), []).append(item)

    for slide_no, slide in enumerate(prs.slides, start=1):
        items = by_slide.get(slide_no, [])
        if not items:
            continue
        for shape in slide.shapes:
            name = getattr(shape, "name", "unnamed")
            for item in items:
                if item.get("shape") and item["shape"] != name:
                    continue
                if not shape.has_text_frame:
                    continue
                para_idx = int(item.get("paragraph", 0))
                paragraphs = shape.text_frame.paragraphs
                if para_idx >= len(paragraphs):
                    continue
                paragraph = paragraphs[para_idx]
                old_full = paragraph.text
                hint = item.get("old_hint")
                if hint and hint not in old_full:
                    changes.append({
                        "slide": slide_no, "shape": name,
                        "old": old_full, "new": item.get("new"),
                        "warning": "old_hint 未命中，已跳过",
                    })
                    continue
                new_full = old_full.replace(hint, item["new"]) if hint else item["new"]
                if new_full == old_full:
                    continue
                if paragraph.runs:
                    paragraph.runs[0].text = new_full
                    for run in paragraph.runs[1:]:
                        run.text = ""
                changes.append({
                    "slide": slide_no, "shape": name,
                    "old": old_full, "new": new_full, "unmatched": False,
                })


def apply_charts(prs, charts_data, changes, unmatched=None):
    used = set()
    for slide_no, slide in enumerate(prs.slides, start=1):
        for shape in slide.shapes:
            if not (getattr(shape, "has_chart", False) and shape.has_chart):
                continue
            key = "slide%d/%s" % (slide_no, getattr(shape, "name", "unnamed"))
            payload = charts_data.get(key)
            if not payload:
                continue
            used.add(key)
            chart = shape.chart
            old_desc = "categories=%s; %s" % (
                [str(c) for c in chart.plots[0].categories],
                {s.name: [v for v in s.values] for s in chart.series})
            chart_data = CategoryChartData()
            chart_data.categories = payload["categories"]
            for series_name, values in payload["series"].items():
                chart_data.add_series(series_name, values)
            chart.replace_data(chart_data)
            changes.append({
                "slide": slide_no, "shape": getattr(shape, "name", "unnamed"),
                "old": old_desc,
                "new": "categories=%s; %s" % (payload["categories"], payload["series"]),
                "kind": "chart",
            })
    if unmatched is not None:
        unmatched.extend(sorted(set(charts_data) - used))


def write_report(path, changes, tokens_all, tokens_used, charts_all, inv,
                 unmatched_charts=None):
    lines = ["# PPT 数据刷新变更报告", ""]
    real = [c for c in changes if not c.get("warning")]
    warns = [c for c in changes if c.get("warning")]

    lines.append("## 已替换（%d 处）" % len(real))
    if real:
        lines.append("")
        lines.append("| 页码 | 位置 | 旧值 | 新值 |")
        lines.append("| --- | --- | --- | --- |")
        for c in real:
            lines.append("| P%d | %s | `%s` | `%s` |" % (
                c["slide"], c["shape"], str(c["old"])[:60].replace("|", "\\|"),
                str(c["new"])[:60].replace("|", "\\|")))
    else:
        lines.append("")
        lines.append("（无）")

    lines += ["", "## 告警 / 跳过（%d 处）" % len(warns)]
    for c in warns:
        lines.append("- P%d / %s：%s" % (c["slide"], c["shape"], c["warning"]))
    if not warns:
        lines.append("")
        lines.append("（无）")

    unused = sorted(set(tokens_used) - set(tokens_all))
    lines += ["", "## 未被使用的数据键（%d 个）" % len(unused)]
    for key in unused:
        lines.append("- `%s`" % key)
    if not unused:
        lines.append("")
        lines.append("（无）")

    left = sorted(t for t in tokens_all if t not in tokens_used)
    lines += ["", "## 未匹配的占位符（%d 个，deck 中存在但未提供值）" % len(left)]
    for key in left:
        lines.append("- `%s`" % key)
    if not left:
        lines.append("")
        lines.append("（无）")

    if unmatched_charts:
        lines += ["", "## ⚠️ 未匹配的图表键（%d 个，数据提供了但 deck 中找不到对应图表）"
                  % len(unmatched_charts)]
        for key in unmatched_charts:
            lines.append("- `%s`" % key)
        lines.append("")
        lines.append("> 图表键格式为 `slide{页码}/{形状名}`，形状名可用 `--scan` 获取。")

    if charts_all:
        lines += ["", "## 图表清单"]
        for c in charts_all:
            lines.append("- %s：series=%s" % (c.get("id"), c.get("series")))

    if inv["dataish"]:
        lines += ["", "## 待人工确认的疑似数据点（%d 处）" % len(inv["dataish"])]
        lines.append("")
        lines.append("| 页码 | 位置 | 值 |")
        lines.append("| --- | --- | --- |")
        for d in inv["dataish"][:60]:
            lines.append("| P%d | %s | `%s` |" % (d["slide"], d["shape"], d["value"]))

    lines += ["", "## 后续人工校验清单", "",
              "- [ ] 单位与口径未变（同比 / 环比、含税 / 不含税）",
              "- [ ] 期间标注同步更新（标题、页脚、图表 X 轴）",
              "- [ ] 合计等于分项之和，占比合计 100%",
              "- [ ] 全文搜索旧值，确认无残留（含备注页）",
              "- [ ] 与数据源抽样比对 3 处", ""]

    Path(path).write_text("\n".join(lines), encoding="utf-8")


def main():
    parser = argparse.ArgumentParser(description="刷新 PowerPoint 中的数据")
    parser.add_argument("--deck", required=True)
    parser.add_argument("--data")
    parser.add_argument("--map")
    parser.add_argument("--out")
    parser.add_argument("--dry-run", action="store_true")
    parser.add_argument("--scan", action="store_true")
    parser.add_argument("--report")
    args = parser.parse_args()

    deck = Path(args.deck)
    if not deck.exists():
        print(json.dumps({"ok": False, "error": "文件不存在: %s" % deck}, ensure_ascii=False))
        sys.exit(1)

    prs = Presentation(str(deck))
    inv_raw = scan_deck(prs)
    inv_tokens = [t["token"] for t in inv_raw["tokens"]]

    if args.scan:
        payload = {
            "ok": True,
            "mode": "scan",
            "deck": str(deck),
            "slides": len(prs.slides.__iter__.__self__._sldIdLst),
            "tokens": inv_raw["tokens"],
            "charts": inv_raw["charts"],
            "dataish": inv_raw["dataish"],
        }
        out = args.out or (deck.stem + "_inventory.json")
        Path(out).write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
        print(json.dumps({
            "ok": True, "inventory": out,
            "tokens": len(inv_tokens),
            "charts": len(inv_raw["charts"]),
            "dataish": len(inv_raw["dataish"]),
        }, ensure_ascii=False, indent=2))
        return

    tokens, charts_data, mapping = {}, {}, []
    if args.data:
        data = json.loads(Path(args.data).read_text(encoding="utf-8"))
        tokens = {k: v for k, v in (data.get("tokens") or {}).items()}
        charts_data = data.get("charts") or {}
    if args.map:
        mapping = json.loads(Path(args.map).read_text(encoding="utf-8"))
        mapping = mapping.get("text", mapping) if isinstance(mapping, dict) else mapping

    if not tokens and not charts_data and not mapping:
        print(json.dumps({"ok": False, "error": "请提供 --data 或 --map"}, ensure_ascii=False))
        sys.exit(1)

    changes = []
    unmatched_charts = []
    apply_tokens(prs, tokens, changes)
    if mapping:
        apply_mapping(prs, mapping, changes)
    if charts_data:
        apply_charts(prs, charts_data, changes, unmatched_charts)

    out_path = args.out or (deck.stem + "_refreshed" + deck.suffix)
    if not args.dry_run:
        if Path(out_path) == deck:
            print(json.dumps({"ok": False, "error": "输出文件与源文件相同，已中止"}, ensure_ascii=False))
            sys.exit(1)
        prs.save(str(out_path))

    if args.report:
        write_report(args.report, changes, inv_tokens, list(tokens.keys()),
                     inv_raw["charts"], inv_raw, unmatched_charts)

    print(json.dumps({
        "ok": True,
        "mode": "dry-run" if args.dry_run else "apply",
        "deck": str(deck),
        "output": None if args.dry_run else str(out_path),
        "report": args.report,
        "changes": len([c for c in changes if not c.get("warning")]),
        "warnings": len([c for c in changes if c.get("warning")]),
        "unmatched_tokens": sorted(t for t in inv_tokens if t not in tokens),
        "unused_keys": sorted(set(tokens.keys()) - set(inv_tokens)),
        "unmatched_chart_keys": unmatched_charts,
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
